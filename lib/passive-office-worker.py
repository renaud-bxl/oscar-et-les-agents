"""Bounded, offline Office preview. JSON values only; no HTML or user paths."""
import base64
import importlib.util
import io
import json
import os
import re
import sys

spec = importlib.util.spec_from_file_location('oscar_reference_reader', os.path.join(os.path.dirname(__file__), 'reference-reader.py'))
reader = importlib.util.module_from_spec(spec)
spec.loader.exec_module(reader)
MAX_TEXT, MAX_CELLS, MAX_BLOCKS = 64000, 40000, 1000
warnings, used, cells = [], 0, 0
partial = False


class PreviewError(ValueError):
    pass


def note(value, incomplete=False):
    global partial
    if value not in warnings and len(warnings) < 12:
        warnings.append(value)
    partial = partial or incomplete


def text(value, limit=4000):
    global used
    value = str(value if value is not None else '').replace('\x00', '')
    maximum = min(limit, max(0, MAX_TEXT - used))
    if len(value) > maximum:
        note('L’aperçu est partiel : le texte dépasse la limite de lecture.', True)
        value = value[:maximum]
    used += len(value)
    return value


def color(value):
    try:
        rgb = value.rgb if value and value.type == 'rgb' else None
        return rgb[-6:].upper() if isinstance(rgb, str) and re.fullmatch('[0-9a-fA-F]{6,8}', rgb) else None
    except (AttributeError, TypeError):
        return None


def word(data):
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    document = Document(io.BytesIO(data))
    blocks = []
    def read_container(container, section=None):
        for item in container.iter_inner_content():
            if len(blocks) >= MAX_BLOCKS or used >= MAX_TEXT:
                note('L’aperçu est partiel : le document dépasse les limites de lecture.', True)
                return
            if isinstance(item, Paragraph):
                if not item.text.strip():
                    continue
                runs = []
                # Hyperlinks contribute plain text, without a clickable target.
                for part in item.iter_inner_content():
                    for run in getattr(part, 'runs', [part]):
                        if len(runs) >= 10000:
                            note('L’aperçu est partiel : le document dépasse les limites de lecture.', True)
                            break
                        value = text(getattr(run, 'text', ''))
                        if not value:
                            continue
                        style = {'bold': bool(run.bold), 'italic': bool(run.italic), 'underline': bool(run.underline)}
                        rgb = str(run.font.color.rgb) if run.font.color.rgb is not None else None
                        if rgb is not None and not re.fullmatch('[0-9a-fA-F]{6}', rgb):
                            rgb = None
                        if rgb:
                            style['color'] = rgb
                        if run.font.size:
                            style['size'] = min(40, max(8, round(run.font.size.pt, 1)))
                        runs.append({'text': value, **style})
                style_name = item.style.name if item.style else ''
                heading = re.match(r'(?:Heading|Titre) ([1-6])$', style_name, re.I)
                block = {'kind': 'paragraph', 'runs': runs, 'heading': int(heading[1]) if heading else 0, 'alignment': {0: 'left', 1: 'center', 2: 'right', 3: 'justify'}.get(item.alignment, 'left')}
                if section:
                    block['section'] = section
                blocks.append(block)
            elif isinstance(item, Table):
                rows = []
                for index, row in enumerate(item.rows):
                    if index >= 200 or used >= MAX_TEXT:
                        note('L’aperçu est partiel : un tableau dépasse les limites de lecture.', True)
                        break
                    rows.append([text(cell.text) for cell in row.cells[:40]])
                    if len(row.cells) > 40:
                        note('L’aperçu est partiel : un tableau dépasse les limites de lecture.', True)
                blocks.append({'kind': 'table', 'rows': rows, **({'section': section} if section else {})})
    seen = set()
    for section in document.sections:
        for label, container in [('header', section.header)]:
            if container.is_linked_to_previous or str(container.part.partname) in seen:
                continue
            seen.add(str(container.part.partname)); read_container(container, label)
    read_container(document)
    for section in document.sections:
        for label, container in [('footer', section.footer)]:
            if container.is_linked_to_previous or str(container.part.partname) in seen:
                continue
            seen.add(str(container.part.partname)); read_container(container, label)
    if document.element.xpath('.//w:drawing|.//w:pict|.//w:txbxContent|.//w:ins|.//w:del|.//w:tbl//w:tbl'):
        note('Les images, zones de texte, tableaux imbriqués et modifications suivies restent dans le document original.', True)
    note('La mise en page affichée est une lecture simplifiée. Le fichier Word original conserve ses pages et ses éléments.')
    return {'kind': 'word', 'blocks': blocks}


def excel(data):
    from openpyxl import load_workbook
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    cached = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    sheets = []
    try:
        if len(workbook.worksheets) > 20:
            note('L’aperçu est limité aux vingt premières feuilles.', True)
        for index, sheet in enumerate(workbook.worksheets[:20]):
            total_rows, total_columns = sheet.max_row or 0, sheet.max_column or 0
            rows = []
            if total_rows > 200 or total_columns > 40:
                note('Chaque feuille affiche au maximum 200 lignes et 40 colonnes.', True)
            cached_rows = cached.worksheets[index].iter_rows(max_row=min(total_rows, 200), max_col=min(total_columns, 40))
            for row, stored_row in zip(sheet.iter_rows(max_row=min(total_rows, 200), max_col=min(total_columns, 40)), cached_rows):
                values = []
                if cells >= MAX_CELLS or used >= MAX_TEXT:
                    note('L’aperçu est partiel : le classeur dépasse les limites de lecture.', True)
                    break
                for column, (cell, stored) in enumerate(zip(row, stored_row), 1):
                    if cell.value is None:
                        continue
                    if cells >= MAX_CELLS or used >= MAX_TEXT:
                        note('L’aperçu est partiel : le classeur dépasse les limites de lecture.', True)
                        break
                    global_cells()
                    formula = cell.data_type == 'f'
                    value = getattr(cell.value, 'text', None) if formula and not isinstance(cell.value, str) else cell.value
                    if value is None and formula:
                        value = '[Formule non prise en charge]'
                        note('Une formule avancée reste consultable dans le fichier Excel original.', True)
                    item = {'column': column, 'address': cell.coordinate, 'text': text(value), 'formula': formula}
                    if formula and stored.value is not None:
                        item['cachedValue'] = text(stored.value)
                    if cell.font:
                        item.update(bold=bool(cell.font.bold), italic=bool(cell.font.italic))
                        rgb = color(cell.font.color)
                        if rgb:
                            item['color'] = rgb
                    if cell.fill and cell.fill.patternType == 'solid':
                        rgb = color(cell.fill.fgColor)
                        if rgb:
                            item['background'] = rgb
                    if isinstance(cell.number_format, str) and cell.number_format not in ('General', '@'):
                        item['numberFormat'] = cell.number_format[:120]
                    values.append(item)
                if values:
                    rows.append({'number': int(values[0]['address'].lstrip('ABCDEFGHIJKLMNOPQRSTUVWXYZ')), 'cells': values})
            sheets.append({'name': sheet.title, 'state': sheet.sheet_state, 'totalRows': total_rows, 'totalColumns': total_columns, 'rows': rows})
        note('Les formules ne sont jamais calculées. Une valeur enregistrée peut dater du dernier calcul dans Excel.')
        note('Les graphiques, images, fusions et mises en page avancées restent dans le fichier Excel original.')
    finally:
        workbook.close(); cached.close()
    return {'kind': 'spreadsheet', 'sheets': sheets}


def global_cells():
    global cells
    cells += 1


def main():
    reader.protect_process()
    raw = sys.stdin.buffer.read(8 * 1024 * 1024 + 1)
    if len(raw) > 8 * 1024 * 1024:
        raise PreviewError('Le document à lire dépasse les limites prévues.')
    job = json.loads(raw)
    if not isinstance(job, dict) or set(job) != {'type', 'base64'} or job['type'] not in ('docx', 'xlsx'):
        raise PreviewError('Ce lecteur accepte uniquement les fichiers Word et Excel.')
    data = base64.b64decode(job['base64'], validate=True)
    if not data or len(data) > reader.FILE_LIMIT:
        raise PreviewError('Chaque fichier doit faire au maximum 5 Mo.')
    parts = reader.office_parts(data)
    if job['type'] == 'docx' and any(name in parts for name in ('word/footnotes.xml', 'word/endnotes.xml')):
        note('Les notes de bas de page et de fin restent dans le document Word original.', True)
    result = word(data) if job['type'] == 'docx' else excel(data)
    result.update(status='partial' if partial else 'ready', warnings=warnings)
    print(json.dumps(result, ensure_ascii=False, allow_nan=False))


if __name__ == '__main__':
    try:
        main()
    except Exception as error:
        message = str(error) if isinstance(error, PreviewError) else 'Ce document n’a pas pu être affiché. Ouvrez son fichier original pour le vérifier.'
        print(json.dumps({'error': message}, ensure_ascii=False)); sys.exit(1)
