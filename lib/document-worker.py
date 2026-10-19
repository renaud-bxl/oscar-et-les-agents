"""Local document conversion. JSON in/out; no user paths, network or shell calls."""
import base64
import copy
import csv
import datetime
import io
import json
import os
import re
import sys
import zipfile
from xml.sax.saxutils import escape

LIMIT = 5 * 1024 * 1024
FIXED = datetime.datetime(2020, 1, 1, 12)
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
S = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
P = 'http://schemas.openxmlformats.org/package/2006/relationships'
warnings = []


def note(message):
    if message not in warnings:
        warnings.append(message)


def xml_parse(data):
    from lxml import etree
    if b'<!DOCTYPE' in data.upper() or b'<!ENTITY' in data.upper():
        raise ValueError('Les déclarations XML externes ne sont pas acceptées.')
    return etree.fromstring(data, etree.XMLParser(resolve_entities=False, no_network=True, huge_tree=False))


def xml_bytes(root):
    from lxml import etree
    return etree.tostring(root, encoding='UTF-8', xml_declaration=True, standalone=True)


def zip_parts(data):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        if len(archive.infolist()) > 2000 or sum(item.file_size for item in archive.infolist()) > 40 * 1024 * 1024:
            raise ValueError('Le document compressé dépasse les limites de sécurité.')
        result = {}
        for item in archive.infolist():
            if item.filename.startswith('/') or '..' in item.filename.split('/') or '\\' in item.filename or item.external_attr >> 16 & 0o170000 == 0o120000:
                raise ValueError('Les chemins dangereux et les liens ne sont pas acceptés dans un document.')
            if item.file_size > 10 * 1024 * 1024 or item.file_size > max(1024 * 1024, item.compress_size * 200):
                raise ValueError('Le document compressé dépasse les limites de sécurité.')
            result[item.filename] = archive.read(item)
        return result


def pack(parts):
    output = io.BytesIO()
    with zipfile.ZipFile(output, 'w', compression=zipfile.ZIP_DEFLATED, compresslevel=6) as archive:
        for name in sorted(parts):
            item = zipfile.ZipInfo(name, (2020, 1, 1, 12, 0, 0))
            item.compress_type = zipfile.ZIP_DEFLATED
            item.external_attr = 0o600 << 16
            archive.writestr(item, parts[name])
    return output.getvalue()


def stable_zip(data):
    parts = zip_parts(data)
    if 'docProps/core.xml' in parts:
        root = xml_parse(parts['docProps/core.xml'])
        for name in ['created', 'modified']:
            element = root.find('{http://purl.org/dc/terms/}' + name)
            if element is not None:
                element.text = '2020-01-01T12:00:00Z'
        parts['docProps/core.xml'] = xml_bytes(root)
    return pack(parts)


def plain(value):
    value = re.sub(r'!\[([^]]*)\]\([^)]*\)', r'[Image : \1]', value)
    return re.sub(r'\[([^]]+)\]\(([^)]+)\)', r'\1 (\2)', value)


def blocks(text):
    lines = text.replace('\r\n', '\n').split('\n')
    result, index = [], 0
    while index < len(lines):
        line = lines[index]
        if line.startswith('```'):
            code = []
            index += 1
            while index < len(lines) and not lines[index].startswith('```'):
                code.append(lines[index]); index += 1
            result.append(('code', '\n'.join(code)))
        elif index + 1 < len(lines) and '|' in line and re.match(r'^\s*\|?\s*:?-{3,}', lines[index + 1]):
            rows = [[cell.strip() for cell in line.strip().strip('|').split('|')]]
            index += 2
            while index < len(lines) and '|' in lines[index] and lines[index].strip():
                rows.append([cell.strip() for cell in lines[index].strip().strip('|').split('|')]); index += 1
            result.append(('table', rows)); continue
        elif line.strip():
            heading = re.match(r'^(#{1,6})\s+(.*)', line)
            if heading:
                result.append(('h' + str(min(3, len(heading[1]))), plain(heading[2])))
            elif re.match(r'^\s*[-*+]\s+', line):
                result.append(('bullet', plain(re.sub(r'^\s*[-*+]\s+', '', line))))
            elif re.match(r'^\s*\d+[.)]\s+', line):
                result.append(('number', plain(re.sub(r'^\s*\d+[.)]\s+', '', line))))
            else:
                result.append(('paragraph', plain(line)))
        index += 1
    return result


def rows_for(job, parsed):
    text, source = job['text'], job['sourceType']
    if source == 'csv':
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;\t')
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.reader(io.StringIO(text), dialect))
    elif source == 'json':
        data = json.loads(text)
        if isinstance(data, list) and data and all(isinstance(row, dict) for row in data):
            keys = list(dict.fromkeys(key for row in data for key in row))
            rows = [keys] + [[row.get(key, '') for key in keys] for row in data]
        elif isinstance(data, list):
            rows = [['Valeur']] + [[value] for value in data]
        elif isinstance(data, dict):
            rows = [['Clé', 'Valeur']] + [[key, value] for key, value in data.items()]
        else:
            rows = [['Valeur'], [data]]
    else:
        rows = next((value for kind, value in parsed if kind == 'table'), None)
        if rows is None:
            rows = [['Texte']] + [[plain(line).strip('# ')] for line in text.splitlines() if line.strip()]
    if not rows or len(rows) > 2001 or max(map(len, rows)) > 40 or sum(map(len, rows)) > 40000:
        raise ValueError('Le tableau dépasse la limite de 2 000 lignes et 40 colonnes.')
    return [[json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value for value in row] for row in rows]


def inline_runs(paragraph, value):
    for token in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', value):
        if token.startswith('**') and token.endswith('**'):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith('`') and token.endswith('`'):
            run = paragraph.add_run(token[1:-1]); run.font.name = 'Courier New'
        else:
            paragraph.add_run(token)


def word_new(job, parsed):
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    document = Document()
    document.core_properties.created = FIXED
    document.core_properties.modified = FIXED
    document.core_properties.author = 'Oscar et les Agents'
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2), Cm(1.8)
    section.left_margin = section.right_margin = Cm(2)
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3']:
        style = document.styles[name]
        style.font.name = 'Calibri'
        style.font.color.rgb = RGBColor(0, 0, 0)
    document.styles['Normal'].font.size = Pt(11)
    document.styles['Normal'].paragraph_format.space_after = Pt(7)
    for border in document.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    if not parsed or not (parsed[0][0].startswith('h') and parsed[0][1] == job['title']):
        document.add_paragraph(job['title'], 'Title')
    for index, (kind, value) in enumerate(parsed):
        if kind == 'table':
            columns = max(map(len, value))
            if columns > 12 or len(value) > 500:
                raise ValueError('Le tableau est trop grand pour une page Word ou PDF.')
            table = document.add_table(rows=0, cols=columns)
            table.style = 'Table Grid'
            for index, row in enumerate(value):
                cells = table.add_row().cells
                for column, text in enumerate(row):
                    cells[column].text = str(text)
                    if index == 0:
                        for run in cells[column].paragraphs[0].runs:
                            run.bold = True
                        fill = OxmlElement('w:shd'); fill.set(qn('w:fill'), 'E8EEE4'); cells[column]._tc.get_or_add_tcPr().append(fill)
        else:
            style = 'Title' if index == 0 and kind.startswith('h') and value == job['title'] else {'h1': 'Heading 1', 'h2': 'Heading 2', 'h3': 'Heading 3', 'bullet': 'List Bullet', 'number': 'List Number'}.get(kind)
            paragraph = document.add_paragraph(style=style)
            inline_runs(paragraph, value)
            if kind == 'code':
                for run in paragraph.runs:
                    run.font.name = 'Courier New'; run.font.size = Pt(9)
    output = io.BytesIO(); document.save(output)
    return stable_zip(output.getvalue())


def word_template(job, parsed, template):
    from lxml import etree
    parts = zip_parts(template)
    root = xml_parse(parts['word/document.xml'])
    lines = []
    for kind, value in parsed:
        if kind == 'table':
            lines.extend(' | '.join(map(str, row)) for row in value)
        else:
            prefix = '• ' if kind == 'bullet' else ''
            lines.extend((prefix + re.sub(r'\*\*([^*]+)\*\*|`([^`]+)`', lambda match: match[1] or match[2], value)).splitlines())
    content = '\n'.join(lines).strip()
    matches = []
    for paragraph in root.iter('{' + W + '}p'):
        texts = list(paragraph.iter('{' + W + '}t'))
        combined = ''.join(item.text or '' for item in texts)
        if '{{CONTENU}}' in combined:
            matches.append((paragraph, combined, texts))
    for paragraph, combined, texts in matches:
        if combined.strip() == '{{CONTENU}}':
            parent = paragraph.getparent(); offset = parent.index(paragraph)
            paragraph_properties = paragraph.find('{' + W + '}pPr')
            run_properties = next(iter(paragraph.iter('{' + W + '}rPr')), None)
            generated = xml_parse(zip_parts(word_new(job, parsed))['word/document.xml']).find('{' + W + '}body')
            elements = [element for element in generated if element.tag != '{' + W + '}sectPr']
            if not parsed or not (parsed[0][0].startswith('h') and parsed[0][1] == job['title']):
                elements = elements[1:]  # The model already owns its title block.
            section = next(reversed(list(root.iter('{' + W + '}sectPr'))), None)
            page = section.find('{' + W + '}pgSz') if section is not None else None
            margins = section.find('{' + W + '}pgMar') if section is not None else None
            width = int(page.get('{' + W + '}w', '11906')) if page is not None else 11906
            if margins is not None:
                width -= int(margins.get('{' + W + '}left', '1440')) + int(margins.get('{' + W + '}right', '1440'))
            else:
                width -= 2880
            # A content marker may be inside a narrow table cell. The generated
            # table must fit that cell rather than the entire page width.
            container = paragraph.getparent()
            while container is not None and container.tag != '{' + W + '}tc':
                container = container.getparent()
            if container is not None:
                cell_width = container.find('{' + W + '}tcPr/' + '{' + W + '}tcW')
                if cell_width is not None and cell_width.get('{' + W + '}type') == 'dxa':
                    width = min(width, max(300, int(cell_width.get('{' + W + '}w', str(width)))))
            for index, element in enumerate(elements):
                replacement = copy.deepcopy(element)
                if replacement.tag == '{' + W + '}p' and paragraph_properties is not None:
                    properties = replacement.find('{' + W + '}pPr')
                    if properties is None:
                        replacement.insert(0, copy.deepcopy(paragraph_properties))
                    else:
                        for property in paragraph_properties:
                            if properties.find(property.tag) is None:
                                properties.append(copy.deepcopy(property))
                for run in replacement.iter('{' + W + '}r'):
                    if run_properties is None:
                        continue
                    properties = run.find('{' + W + '}rPr')
                    if properties is None:
                        run.insert(0, copy.deepcopy(run_properties))
                    else:
                        for property in run_properties:
                            if properties.find(property.tag) is None:
                                properties.append(copy.deepcopy(property))
                if replacement.tag == '{' + W + '}tbl':
                    table_width = replacement.find('{' + W + '}tblPr/' + '{' + W + '}tblW')
                    if table_width is not None:
                        table_width.set('{' + W + '}w', str(width))
                        table_width.set('{' + W + '}type', 'dxa')
                    columns = list(replacement.iter('{' + W + '}gridCol'))
                    if columns:
                        column_width = max(300, width // len(columns))
                        for column in columns:
                            column.set('{' + W + '}w', str(column_width))
                        for cell_width in replacement.iter('{' + W + '}tcW'):
                            cell_width.set('{' + W + '}w', str(column_width))
                parent.insert(offset + index, replacement)
            parent.remove(paragraph)
            if parent.tag == '{' + W + '}tc' and (not len(parent) or parent[-1].tag != '{' + W + '}p'):
                etree.SubElement(parent, '{' + W + '}p')
        else:
            # Keep the surrounding runs intact, including split placeholders.
            start = combined.index('{{CONTENU}}'); end = start + len('{{CONTENU}}'); cursor = 0
            for element in texts:
                original = element.text or ''; left, right = cursor, cursor + len(original); cursor = right
                if right <= start or left >= end:
                    continue
                before = original[:max(0, start - left)]
                after = original[max(0, end - left):] if end < right else ''
                element.text = before + (content if left <= start < right else '') + after
            note('La zone intégrée dans une phrase reprend la mise en forme du premier segment du modèle.')
    if not matches:
        generated = xml_parse(zip_parts(word_new(job, parsed))['word/document.xml'])
        body = root.find('{' + W + '}body'); source = generated.find('{' + W + '}body')
        position = len(body) - (1 if len(body) and body[-1].tag == '{' + W + '}sectPr' else 0)
        for element in source:
            if element.tag != '{' + W + '}sectPr':
                body.insert(position, copy.deepcopy(element)); position += 1
        note('Ce modèle n’a pas de zone {{CONTENU}} : le nouveau texte a été ajouté à la fin.')
    parts['word/document.xml'] = xml_bytes(root)
    return pack(parts)


def scalar(value):
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value if len(str(value).replace('-', '').replace('.', '')) <= 15 else str(value)
    # Preserve identifiers such as 00123 and numbers Excel cannot store exactly.
    if isinstance(value, str) and re.fullmatch(r'-?(?:0|[1-9]\d{0,11})(?:\.\d{1,8})?', value.strip()) and len(re.sub(r'\D', '', value)) <= 15:
        return float(value) if '.' in value else int(value)
    return str(value if value is not None else '')[:32767]


def spreadsheet_new(rows):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    workbook = Workbook(); sheet = workbook.active; sheet.title = 'Livraison Oscar'
    workbook.properties.creator = 'Oscar et les Agents'; workbook.properties.created = FIXED; workbook.properties.modified = FIXED
    sheet.freeze_panes = 'A2'
    for index, row in enumerate(rows, start=1):
        for column, value in enumerate(row, start=1):
            cell = sheet.cell(index, column); value = scalar(value); cell.value = value
            if isinstance(value, str):
                cell.data_type = 's'  # Text from a model/CSV must never become a formula.
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            cell.font = Font(name='Calibri', size=11, bold=index == 1, color='203429' if index == 1 else '27362C')
            if index == 1:
                cell.fill = PatternFill('solid', fgColor='E6EEDC'); cell.border = Border(bottom=Side(style='thin', color='A9BC99'))
            elif index % 2 == 0:
                cell.fill = PatternFill('solid', fgColor='F6F8F2')
    for column in range(1, max(map(len, rows)) + 1):
        values = [str(row[column - 1]) for row in rows if len(row) >= column]
        sheet.column_dimensions[get_column_letter(column)].width = min(58, max(16, min(46, max(map(len, values), default=16)) + 2))
    sheet.auto_filter.ref = sheet.dimensions
    sheet.print_options.horizontalCentered = True
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'; sheet.page_setup.paperSize = sheet.PAPERSIZE_A4; sheet.page_setup.fitToWidth = 1; sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = '1:1'
    output = io.BytesIO(); workbook.save(output)
    return stable_zip(output.getvalue())


def spreadsheet_template(rows, template):
    from lxml import etree
    parts = zip_parts(template)
    source = zip_parts(spreadsheet_new(rows))
    workbook = xml_parse(parts['xl/workbook.xml']); relations = xml_parse(parts['xl/_rels/workbook.xml.rels']); types = xml_parse(parts['[Content_Types].xml'])
    sheets = workbook.find('{' + S + '}sheets'); names = {item.get('name') for item in sheets}
    header_style, body_style = '0', '0'
    if len(sheets):
        source_id = sheets[0].get('{' + R + '}id')
        relation = next((item for item in relations if item.get('Id') == source_id), None)
        if relation is not None:
            target = relation.get('Target', '')
            target = target.lstrip('/') if target.startswith('/') else 'xl/' + target
            if target in parts:
                original = xml_parse(parts[target]).find('{' + S + '}sheetData')
                if original is not None and len(original):
                    header_style = next((cell.get('s', '0') for cell in original[0] if cell.tag == '{' + S + '}c'), '0')
                    if len(original) > 1:
                        body_style = next((cell.get('s', '0') for cell in original[1] if cell.tag == '{' + S + '}c'), '0')
    name, suffix = 'Livraison Oscar', 2
    while name in names:
        name = 'Livraison Oscar ' + str(suffix); suffix += 1
    number = 1
    while 'xl/worksheets/oscar' + str(number) + '.xml' in parts:
        number += 1
    ids = {item.get('Id') for item in relations}; relation_id = 'rIdOscar' + str(number)
    while relation_id in ids:
        number += 1; relation_id = 'rIdOscar' + str(number)
    filename = 'worksheets/oscar' + str(number) + '.xml'
    # Inline strings use no shared-string/style indexes from another workbook.
    worksheet = etree.Element('{' + S + '}worksheet', nsmap={None: S})
    views = etree.SubElement(worksheet, '{' + S + '}sheetViews')
    etree.SubElement(views, '{' + S + '}sheetView', workbookViewId='0', showGridLines='1')
    cols = etree.SubElement(worksheet, '{' + S + '}cols')
    for column in range(1, max(map(len, rows)) + 1):
        longest = max((len(str(row[column - 1])) for row in rows if len(row) >= column), default=14)
        etree.SubElement(cols, '{' + S + '}col', min=str(column), max=str(column), width=str(min(58, max(16, longest + 3))), customWidth='1')
    data = etree.SubElement(worksheet, '{' + S + '}sheetData')
    from openpyxl.utils import get_column_letter
    for index, row in enumerate(rows, start=1):
        output = etree.SubElement(data, '{' + S + '}row', r=str(index), ht='24' if index == 1 else '21', customHeight='1')
        for column, value in enumerate(row, start=1):
            value = scalar(value); cell = etree.SubElement(output, '{' + S + '}c', r=get_column_letter(column) + str(index))
            cell.set('s', header_style if index == 1 else body_style if index % 2 == 0 else '0')
            if isinstance(value, (int, float)):
                etree.SubElement(cell, '{' + S + '}v').text = str(value)
            else:
                cell.set('t', 'inlineStr'); string = etree.SubElement(cell, '{' + S + '}is'); text = etree.SubElement(string, '{' + S + '}t'); text.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve'); text.text = value
    etree.SubElement(worksheet, '{' + S + '}printOptions', horizontalCentered='1', gridLines='1')
    etree.SubElement(worksheet, '{' + S + '}pageMargins', left='0.35', right='0.35', top='0.5', bottom='0.5', header='0.2', footer='0.2')
    etree.SubElement(worksheet, '{' + S + '}pageSetup', paperSize='9', orientation='landscape', fitToWidth='1', fitToHeight='0')
    sheet_id = max([int(item.get('sheetId', '0')) for item in sheets] + [0]) + 1
    item = etree.SubElement(sheets, '{' + S + '}sheet', name=name, sheetId=str(sheet_id)); item.set('{' + R + '}id', relation_id)
    etree.SubElement(relations, '{' + P + '}Relationship', Id=relation_id, Type=R + '/worksheet', Target=filename)
    namespace = '{http://schemas.openxmlformats.org/package/2006/content-types}'
    etree.SubElement(types, namespace + 'Override', PartName='/xl/' + filename, ContentType='application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml')
    parts['xl/workbook.xml'] = xml_bytes(workbook); parts['xl/_rels/workbook.xml.rels'] = xml_bytes(relations); parts['[Content_Types].xml'] = xml_bytes(types); parts['xl/' + filename] = xml_bytes(worksheet)
    note('Les feuilles du modèle sont conservées. Les nouvelles données sont placées dans une feuille Livraison Oscar distincte.')
    return pack(parts)


def pdf_new(job, parsed):
    import reportlab
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle
    fontdir = os.path.join(os.path.dirname(reportlab.__file__), 'fonts')
    for name, filename in [('Office', 'Vera.ttf'), ('OfficeBold', 'VeraBd.ttf'), ('OfficeMono', 'Vera.ttf')]:
        with open(os.path.join(fontdir, filename), 'rb') as font:
            pdfmetrics.registerFont(TTFont(name, font))
    pdfmetrics.registerFontFamily('Office', normal='Office', bold='OfficeBold', italic='Office', boldItalic='OfficeBold')
    styles = getSampleStyleSheet()
    for style in styles.byName.values():
        style.fontName = 'Office'; style.textColor = colors.HexColor('#17251D')
    styles['Normal'].fontSize = 10; styles['Normal'].leading = 15; styles['Normal'].spaceAfter = 7
    styles['Title'].fontSize = 23; styles['Title'].leading = 29; styles['Title'].spaceAfter = 18
    for name, size in [('Heading1', 16), ('Heading2', 13), ('Heading3', 11)]:
        styles[name].fontName = 'OfficeBold'; styles[name].fontSize = size; styles[name].leading = size + 5; styles[name].spaceBefore = 12; styles[name].spaceAfter = 7
    styles.add(ParagraphStyle('OfficeCode', parent=styles['Normal'], fontName='OfficeMono', fontSize=8, leading=11, backColor=colors.HexColor('#F2F5EE'), borderPadding=7))
    styles.add(ParagraphStyle('OfficeCell', parent=styles['Normal'], fontSize=8, leading=11, spaceAfter=0, alignment=TA_LEFT))
    glyphs = pdfmetrics.getFont('Office').face.charToGlyph
    def markup(value):
        value = str(value)
        missing = {character for character in value if ord(character) > 31 and ord(character) not in glyphs}
        if missing:
            note('Certains symboles sans glyphe dans la police PDF sont indiqués par leur code Unicode.')
            value = ''.join('[U+%04X]' % ord(character) if character in missing else character for character in value)
        value = escape(value)
        value = re.sub(r'\*\*([^*]+)\*\*', r'<b>\1</b>', value)
        return value.replace('\n', '<br/>')
    story = []
    if not parsed or not (parsed[0][0].startswith('h') and parsed[0][1] == job['title']):
        story.append(Paragraph(markup(job['title']), styles['Title']))
    for kind, value in parsed:
        if kind == 'table':
            columns = max(map(len, value))
            if columns > 12 or len(value) > 500:
                raise ValueError('Le tableau est trop grand pour une page Word ou PDF.')
            rows = [[Paragraph(markup(str(cell)[:2000]), styles['OfficeCell']) for cell in row] + [''] * (columns - len(row)) for row in value]
            table = LongTable(rows, colWidths=[481 / columns] * columns, repeatRows=1, hAlign='LEFT')
            table.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#E6EEDC')), ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('LEFTPADDING', (0, 0), (-1, -1), 7), ('RIGHTPADDING', (0, 0), (-1, -1), 7), ('TOPPADDING', (0, 0), (-1, -1), 7), ('BOTTOMPADDING', (0, 0), (-1, -1), 7), ('LINEBELOW', (0, 0), (-1, 0), .6, colors.HexColor('#A9BC99')), ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F7F9F3')])]))
            story += [table, Spacer(1, 12)]
        else:
            style = styles[{'h1': 'Heading1', 'h2': 'Heading2', 'h3': 'Heading3', 'code': 'OfficeCode'}.get(kind, 'Normal')]
            prefix = '• ' if kind == 'bullet' else '– ' if kind == 'number' else ''
            story.append(Paragraph(markup(prefix + value), style))
    def footer(canvas, document):
        canvas.saveState(); canvas.setFont('Office', 8); canvas.setFillColor(colors.HexColor('#77856F')); canvas.drawString(57, 29, 'Oscar et les Agents'); canvas.drawRightString(A4[0] - 57, 29, str(document.page)); canvas.restoreState()
    output = io.BytesIO()
    document = SimpleDocTemplate(output, pagesize=A4, rightMargin=57, leftMargin=57, topMargin=50, bottomMargin=47, title=job['title'], author='Oscar et les Agents', invariant=1)
    document.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def main():
    raw = sys.stdin.buffer.read(8 * 1024 * 1024 + 1)
    if len(raw) > 8 * 1024 * 1024:
        raise ValueError('Le contenu à exporter est trop volumineux.')
    job = json.loads(raw)
    if not isinstance(job.get('text'), str) or len(job['text'].encode('utf-8')) > 512 * 1024:
        raise ValueError('Le texte à exporter est vide ou dépasse 512 Ko.')
    if not isinstance(job.get('formats'), list) or not 1 <= len(job['formats']) <= 3 or len(set(job['formats'])) != len(job['formats']) or any(value not in ['docx', 'pdf', 'xlsx'] for value in job['formats']):
        raise ValueError('Choisissez Word, PDF ou Excel.')
    parsed = blocks(job['text']); rows = rows_for(job, parsed)
    if job['sourceType'] in ['csv', 'json']:
        parsed = [('table', rows)]
    template = job.get('template'); template_bytes = base64.b64decode(template['base64'], validate=True) if template else None
    if template_bytes and len(template_bytes) > LIMIT:
        raise ValueError('Chaque fichier doit faire au maximum 5 Mo.')
    if template and template['type'] == 'pdf':
        note('Le PDF de référence sert d’exemple ; sa mise en page n’est pas reproduite automatiquement.')
    files = []
    for format in job['formats']:
        if format == 'docx':
            output = word_template(job, parsed, template_bytes) if template and template['type'] == 'docx' else word_new(job, parsed)
        elif format == 'xlsx':
            output = spreadsheet_template(rows, template_bytes) if template and template['type'] == 'xlsx' else spreadsheet_new(rows)
        else:
            output = pdf_new(job, parsed)
        if len(output) > LIMIT:
            raise ValueError('Le document généré dépasse la limite de 5 Mo.')
        files.append({'format': format, 'base64': base64.b64encode(output).decode('ascii')})
    print(json.dumps({'files': files, 'warnings': warnings}, ensure_ascii=False))


if __name__ == '__main__':
    try:
        main()
    except Exception as error:
        message = str(error) if isinstance(error, ValueError) else 'La génération du document n’a pas abouti. Vérifiez le contenu et le modèle.'
        print(json.dumps({'error': message}, ensure_ascii=False))
        sys.exit(1)
