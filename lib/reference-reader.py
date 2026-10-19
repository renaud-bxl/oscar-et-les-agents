"""Bounded local reference reading: JSON stdin/stdout; no paths or network."""
import base64
import io
import json
import os
import re
import sys
import threading
import time
import zipfile

FILE_LIMIT = 5 * 1024 * 1024
TEXT_LIMIT = 64000
MEMORY_LIMIT = 768 * 1024 * 1024
warnings = []
partial = False


def protect_process():
    try:
        import resource
        resource.setrlimit(resource.RLIMIT_CPU, (12, 14))
        if sys.platform.startswith('linux'):
            resource.setrlimit(resource.RLIMIT_AS, (MEMORY_LIMIT, MEMORY_LIMIT))
        # Darwin does not reliably allow RLIMIT_AS/DATA. Its peak-RSS monitor
        # complements bounded streams and the parent process' 15-second timer.
        def monitor():
            while True:
                used = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
                if sys.platform != 'darwin':
                    used *= 1024
                if used > MEMORY_LIMIT:
                    os.write(1, json.dumps({'error': 'La lecture du document a dépassé la mémoire prévue.'}, ensure_ascii=False).encode('utf-8'))
                    os._exit(1)
                time.sleep(0.05)
        threading.Thread(target=monitor, daemon=True).start()
    except (ImportError, OSError, ValueError):
        pass
    import socket
    def offline(*args, **kwargs):
        raise OSError('Network disabled for reference reading')
    socket.socket = offline
    socket.create_connection = offline


def note(message, incomplete=False):
    global partial
    if message not in warnings and len(warnings) < 12:
        warnings.append(message)
    partial = partial or incomplete


class TextFull(Exception):
    pass


class Text:
    def __init__(self, limit):
        self.parts, self.length, self.limit = [], 0, limit
    def add(self, value):
        value = str(value).replace('\x00', '').strip()
        if not value:
            return
        remaining = self.limit - self.length - (1 if self.parts else 0)
        if remaining <= 0 or len(value) > remaining:
            if remaining > 0:
                self.parts.append(value[:remaining]); self.length = self.limit
            note('Le texte dépasse la limite de lecture ; seule la première partie est préparée.', True)
            raise TextFull()
        self.parts.append(value); self.length += len(value) + (1 if len(self.parts) > 1 else 0)
    def value(self):
        return '\n'.join(self.parts)


def office_parts(data):
    result = {}
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        items = archive.infolist()
        if len(items) > 2000 or sum(item.file_size for item in items) > 40 * 1024 * 1024:
            raise ValueError('Le document compressé dépasse les limites de sécurité.')
        for item in items:
            name = item.filename
            if name in result or name.startswith('/') or '..' in name.split('/') or '\\' in name or item.external_attr >> 16 & 0o170000 == 0o120000:
                raise ValueError('Les chemins dangereux et les liens ne sont pas acceptés dans un document.')
            if item.flag_bits & 1 or item.compress_type not in (0, 8) or item.file_size > 10 * 1024 * 1024 or item.file_size > max(1024 * 1024, item.compress_size * 200):
                raise ValueError('Le document compressé dépasse les limites de sécurité.')
            result[name] = archive.read(item)
            if name.lower().endswith(('.xml', '.rels')):
                from lxml import etree
                root = etree.fromstring(result[name], etree.XMLParser(resolve_entities=False, no_network=True, huge_tree=False))
                if root.getroottree().docinfo.doctype:
                    raise ValueError('Les déclarations XML externes ne sont pas acceptées.')
    return result


def read_word(data, text):
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    parts = office_parts(data)
    document = Document(io.BytesIO(data))
    fonts, colors = set(), set()
    namespace = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    for color in document.element.xpath('.//w:color'):
        value = color.get(namespace + 'val')
        if value and re.fullmatch('[0-9a-fA-F]{6}', value): colors.add(value.upper())
    for font in document.element.xpath('.//w:rFonts'):
        for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
            value = font.get(namespace + key)
            if value: fonts.add(value)
    used_styles = set(document.element.xpath('.//w:pStyle/@w:val'))
    for style in list(document.styles)[:300]:
        if style.style_id not in used_styles and style.name != 'Normal': continue
        font = getattr(style, 'font', None)
        if font is None: continue
        if font.name: fonts.add(font.name)
        if font.color.rgb: colors.add(str(font.color.rgb))
    hints = {'fonts': sorted(fonts)[:12], 'colors': sorted(colors)[:12]}
    if document.sections:
        section = document.sections[0]
        hints['pageCentimeters'] = {name: round(getattr(section, name).cm, 2) for name in ('page_width', 'page_height', 'top_margin', 'bottom_margin', 'left_margin', 'right_margin') if getattr(section, name) is not None}
    count = 0
    def container(block, depth=0):
        nonlocal count
        if depth > 8:
            note('Certains tableaux imbriqués dépassent la limite de lecture.', True); return
        for item in block.iter_inner_content():
            count += 1
            if count > 10000:
                note('Le document contient trop de blocs ; la lecture est partielle.', True); return
            if isinstance(item, Paragraph):
                if item.text.strip():
                    style = item.style.name if item.style else ''
                    text.add(f'[{style}] {item.text}' if style and style not in ('Normal', 'Body Text') else item.text)
            elif isinstance(item, Table):
                table_started = False
                for row_index, row in enumerate(item.rows):
                    if row_index >= 2000:
                        note('Un tableau dépasse 2 000 lignes ; sa lecture est partielle.', True); break
                    cells = row.cells
                    if len(cells) > 40: note('Un tableau dépasse 40 colonnes ; sa lecture est partielle.', True)
                    value = '\t'.join(cell.text.replace('\n', ' / ') for cell in cells[:40])
                    if value.strip():
                        if not table_started: text.add('[Tableau]'); table_started = True
                        text.add(value)
    container(document)
    seen = set()
    for section in document.sections:
        for name in ('header', 'first_page_header', 'even_page_header', 'footer', 'first_page_footer', 'even_page_footer'):
            block = getattr(section, name)
            # Reading an undefined header through python-docx would create one.
            if block.is_linked_to_previous: continue
            key = str(block.part.partname)
            if key in seen: continue
            seen.add(key); container(block)
    if any(name.startswith('word/media/') for name in parts): note('Les images et dessins restent consultables dans le document original ; ils ne sont pas décrits dans cet extrait.')
    if any(name in parts for name in ('word/footnotes.xml', 'word/endnotes.xml')): note('Les notes de bas de page et de fin ne sont pas incluses dans cet extrait.', True)
    if document.element.xpath('.//w:txbxContent|.//w:ins|.//w:del'): note('Les zones de texte et les modifications suivies peuvent manquer dans cet extrait.', True)
    if document.element.xpath('.//w:tbl//w:tbl'): note('Les tableaux imbriqués peuvent manquer dans cet extrait.', True)
    if not text.value(): note('Aucun texte détecté dans ce document Word. Les images ne sont pas transcrites.')
    return hints


def read_excel(data, text):
    from openpyxl import load_workbook
    office_parts(data)
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    try:
        if len(workbook.worksheets) > 20: note('Le classeur dépasse 20 feuilles ; sa lecture est partielle.', True)
        for sheet in workbook.worksheets[:20]:
            sheet_started = False
            if (sheet.max_row or 0) > 2000 or (sheet.max_column or 0) > 40: note('Une feuille dépasse 2 000 lignes ou 40 colonnes ; sa lecture est partielle.', True)
            cells_read = 0
            for row in sheet.iter_rows(max_row=min(sheet.max_row or 2000, 2000), max_col=min(sheet.max_column or 40, 40)):
                values = []
                for cell in row:
                    if cell.value is None: continue
                    cells_read += 1
                    if cells_read > 40000: note('Une feuille dépasse 40 000 cellules ; sa lecture est partielle.', True); break
                    value = str(cell.value)
                    if cell.data_type == 'f': value = '[Formule non calculée] ' + value
                    elif cell.number_format not in ('General', '@'): value += f' [Format : {cell.number_format}]'
                    values.append(f'{cell.coordinate} : {value}')
                if values:
                    if not sheet_started:
                        text.add(f'[Feuille : {sheet.title}' + (' — masquée' if sheet.sheet_state != 'visible' else '') + ']'); sheet_started = True
                    text.add(' | '.join(values))
                if cells_read > 40000: break
        note('Les formules sont présentées comme du texte, sans calcul ; les graphiques et images restent dans le fichier original.')
        if not text.value(): note('Aucune cellule renseignée détectée dans ce classeur.')
    finally:
        workbook.close()
    return {}


def read_pdf(data, text):
    import pypdf
    from pypdf.generic import ArrayObject
    limit = 2 * 1024 * 1024
    pypdf.overwrite_configuration(maximum_declared_stream_length=8*1024*1024, array_based_stream_maximum_output_length=limit, jbig2_maximum_output_length=limit, lzw_maximum_output_length=limit, run_length_maximum_output_length=limit, zlib_maximum_output_length=limit, zlib_maximum_recovery_input_length=200000, flate_maximum_columns=40000, flate_maximum_row_length=limit, image_maximum_buffer_size=limit, xmp_maximum_input_length=limit, xmp_maximum_element_count=10000, outline_maximum_entries=2000, outline_maximum_depth=32, page_tree_maximum_entries=2000, page_tree_maximum_depth=32, xform_maximum_invocations_per_extraction=64, jbig2dec_binary=None)
    reader = pypdf.PdfReader(io.BytesIO(data), strict=True, root_object_recovery_limit=2000)
    if reader.is_encrypted: raise ValueError('Ce PDF est protégé par un mot de passe. Fournissez une copie lisible.')
    pages = len(reader.pages)
    if pages > 50: note('Le PDF dépasse 50 pages ; seules les 50 premières sont préparées.', True)
    decoded_total, content_found = 0, False
    for index in range(min(pages, 50)):
        page = reader.pages[index]
        contents = page.get('/Contents')
        if contents is not None:
            contents = contents.get_object()
            streams = contents if isinstance(contents, ArrayObject) else [contents]
            if len(streams) > 64: raise ValueError('Le PDF contient trop de flux pour une lecture sûre.')
            page_bytes = 0
            for stream in streams:
                decoded = stream.get_object().get_data(); page_bytes += len(decoded); decoded_total += len(decoded)
                if page_bytes > limit or decoded_total > 8 * 1024 * 1024: raise ValueError('Les données décompressées du PDF dépassent la limite de lecture.')
        value = page.extract_text() or ''
        if value.strip():
            content_found = True; text.add(f'[Page {index + 1}]'); text.add(value)
    if not content_found: note('Aucun texte détecté dans ce PDF. Il peut s’agir d’un scan ; la reconnaissance de caractères n’est pas activée.')
    else: note('L’ordre de lecture et les tableaux du PDF peuvent différer de la page originale. Les images ne sont pas transcrites.')
    return {'pages': pages, 'pagesRead': min(pages, 50)}


def main():
    protect_process()
    raw = sys.stdin.buffer.read(8 * 1024 * 1024 + 1)
    if len(raw) > 8 * 1024 * 1024: raise ValueError('Le document à lire dépasse les limites prévues.')
    job = json.loads(raw)
    if not isinstance(job, dict) or set(job) - {'type', 'base64', 'maxCharacters'} or job.get('type') not in ('docx', 'xlsx', 'pdf'): raise ValueError('Ce format de référence ne peut pas être lu.')
    maximum = job.get('maxCharacters', TEXT_LIMIT)
    if type(maximum) is not int or not 1 <= maximum <= TEXT_LIMIT: raise ValueError('La limite de lecture du document est invalide.')
    data = base64.b64decode(job['base64'], validate=True)
    if not data or len(data) > FILE_LIMIT: raise ValueError('Chaque fichier doit faire au maximum 5 Mo.')
    text, hints = Text(maximum), {}
    try:
        hints = {'docx': read_word, 'xlsx': read_excel, 'pdf': read_pdf}[job['type']](data, text)
    except TextFull:
        pass
    value = text.value()
    print(json.dumps({'text': value, 'status': 'partial' if partial else 'ready' if value else 'empty', 'warnings': warnings, 'characters': len(value), 'styleHints': hints}, ensure_ascii=False))


if __name__ == '__main__':
    try:
        main()
    except Exception as error:
        message = str(error) if isinstance(error, ValueError) else 'La lecture du document n’a pas abouti. Vérifiez le fichier ou fournissez un exemple texte.'
        print(json.dumps({'error': message}, ensure_ascii=False))
        sys.exit(1)
